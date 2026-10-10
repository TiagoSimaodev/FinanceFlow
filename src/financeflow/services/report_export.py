import csv
import io
from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


EXPORT_COLUMNS = ("Data", "Tipo", "Categoria", "Descrição", "Valor (R$)")


def _safe_spreadsheet_text(value: str) -> str:
    text = value or ""
    if text.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + text
    return text


def export_transactions(rows, file_format: str, month: str) -> tuple[bytes, str, str]:
    if file_format == "csv":
        output = io.StringIO(newline="")
        writer = csv.writer(output, delimiter=";", lineterminator="\r\n")
        writer.writerow(EXPORT_COLUMNS)
        for row in rows:
            writer.writerow(
                (
                    date.fromisoformat(str(row["transaction_date"])).strftime(
                        "%d/%m/%Y"
                    ),
                    "Receita" if row["type"] == "receita" else "Despesa",
                    _safe_spreadsheet_text(row["category_name"]),
                    _safe_spreadsheet_text(row["description"]),
                    f"{Decimal(row['amount_cents']) / 100:.2f}".replace(".", ","),
                )
            )
        return (
            output.getvalue().encode("utf-8-sig"),
            "text/csv; charset=utf-8",
            f"financeflow-movimentacoes-{month}.csv",
        )

    if file_format != "xlsx":
        raise ValueError("Formato de exportação não suportado.")

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Movimentações"
    worksheet.append(EXPORT_COLUMNS)
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="247A55")

    for row in rows:
        worksheet.append(
            (
                date.fromisoformat(str(row["transaction_date"])),
                "Receita" if row["type"] == "receita" else "Despesa",
                _safe_spreadsheet_text(row["category_name"]),
                _safe_spreadsheet_text(row["description"]),
                row["amount_cents"] / 100,
            )
        )
        worksheet.cell(worksheet.max_row, 1).number_format = "dd/mm/yyyy"
        worksheet.cell(worksheet.max_row, 5).number_format = '"R$" #,##0.00'

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    column_widths = (14, 14, 24, 42, 18)
    for index, width in enumerate(column_widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width

    output = io.BytesIO()
    workbook.save(output)
    return (
        output.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        f"financeflow-movimentacoes-{month}.xlsx",
    )
